"""Add every qualifying job from the shortlist to your tracker, with Status New.
Usage: python3 tools/add_to_tracker.py [tracker/my_job_tracker.xlsx] [--dry-run]

- Adds each job that passed your Fit Score gate and isn't in the tracker yet: Date Added, Company, Job Title,
  Location, Work Type, Salary Range, Salary Min, Source, Link, Status New and Fit Score. Resume Match is left
  blank for your assistant (or you) to fill in.
- Flags such as "salary unknown" go in the Assistant Notes column. Your Notes column is never written to.
  If your tracker has no Assistant Notes column, one is added after your last column.
- Saves a backup copy in tracker/backups/ before changing anything.
- Then runs tools/import_tracker.py, so these jobs aren't suggested again.

You review the new rows in the tracker: set Status (Applied, Did Not Apply, ...) and write your reasons in Notes.
Works with .xlsx trackers (needs openpyxl: pip3 install openpyxl). For a CSV tracker, use tools/export_shortlist.py.
"""
import datetime, shutil, subprocess, sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent / "collectors"))
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent / "scoring"))
import common
from common import norm_co, norm_title
from score import shortlist
from export_shortlist import WORK_TYPE

ASSISTANT_NOTES = "Assistant Notes"

def header_row(ws, company, title):
    """Row number and {lowercase header: column} for the first row (of 10) that has the company and title headers."""
    for r in range(1, 11):
        names = {str(c.value).strip().lower(): c.column for c in ws[r] if c.value not in (None, "")}
        if company.lower() in names and title.lower() in names: return r, names
    sys.exit(f"Could not find the '{company}' and '{title}' headers. Check config/tracker_map.json.")

def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry-run" in sys.argv
    path = pathlib.Path(args[0] if args else common.ROOT / "tracker" / "my_job_tracker.xlsx").expanduser()
    if not path.exists(): sys.exit(f"Tracker not found: {path}")
    if path.suffix.lower() != ".xlsx": sys.exit("This tool writes .xlsx trackers. For CSV, run tools/export_shortlist.py.")
    import openpyxl
    tmap = common.load("tracker_map.json"); cols = tmap["columns"]
    wb = openpyxl.load_workbook(path)
    ws = wb[tmap["sheet"]] if tmap.get("sheet") else wb.worksheets[0]
    hr, names = header_row(ws, cols["company"], cols["title"])
    col = lambda name: names.get(name.lower())
    c_co, c_ti = col(cols["company"]), col(cols["title"])
    c_status, c_link = col(cols.get("status", "Status")), col(cols.get("link", "Link"))

    existing, last = set(), hr
    for r in range(hr + 1, ws.max_row + 1):
        co, ti = ws.cell(r, c_co).value, ws.cell(r, c_ti).value
        if co not in (None, "") or ti not in (None, ""):
            last = r
            existing.add(norm_co(str(co or "")) + "|" + norm_title(str(ti or "")))

    jobs, _ = shortlist(common.connect())
    new = [j for j in jobs if norm_co(j["company"]) + "|" + norm_title(j["title"]) not in existing]
    print(f"{len(jobs)} qualifying jobs, {len(jobs) - len(new)} already in the tracker, {len(new)} to add.")
    for j in new: print(f"  + {j['fit']:>3} | {j['company']} | {j['title']}")
    if dry or not new:
        if dry: print("Dry run: tracker not changed.")
        return

    backup = path.parent / "backups" / f"{path.stem}_{datetime.datetime.now():%Y%m%d-%H%M%S}{path.suffix}"
    backup.parent.mkdir(exist_ok=True); shutil.copy(path, backup)

    c_an = col(ASSISTANT_NOTES)
    if not c_an:
        c_an = max(names.values()) + 1
        h = ws.cell(hr, c_an, ASSISTANT_NOTES); h._style = ws.cell(hr, c_an - 1)._style
        ws.column_dimensions[openpyxl.utils.get_column_letter(c_an)].width = 50
    today = datetime.date.today()
    for i, j in enumerate(new, 1):
        r = last + i
        sal = f"${j['smin']:,} - ${j['smax']:,}" if j["smin"] and j["smax"] else (f"${j['smin']:,}" if j["smin"] else None)
        values = {"Date Added": today, "Location": j["location"], "Work Type": WORK_TYPE.get(j.get("group"), None),
                  "Salary Range": sal, "Salary Min": j["smin"], "Source": j["source"], "Fit Score": j["fit"]}
        for name, v in values.items():
            if col(name) and v not in (None, ""): ws.cell(r, col(name)).value = v
        if col("Date Added"): ws.cell(r, col("Date Added")).number_format = "mm/dd/yyyy"
        ws.cell(r, c_co).value = j["company"]; ws.cell(r, c_ti).value = j["title"]
        if c_status: ws.cell(r, c_status).value = "New"
        if c_link:
            ws.cell(r, c_link).value = "Link"; ws.cell(r, c_link).hyperlink = j["url"]
        note = [f"Added by pipeline {today:%m/%d}"] + ([f"Flags: {j['flags']}"] if j.get("flags") else [])
        ws.cell(r, c_an).value = ". ".join(note) + "."
    wb.save(path)
    print(f"Added {len(new)} rows to {path} (backup: {backup.name}).")
    subprocess.run([sys.executable, str(common.ROOT / "tools" / "import_tracker.py"), str(path)], check=False)

if __name__ == "__main__":
    main()
