"""Copy an existing job tracker (yours, in any layout) into this project's tracker template.
Usage: python3 tools/migrate_tracker.py path/to/old_tracker.xlsx [--sheet "Tab name"] [--out tracker/my_job_tracker.xlsx] [--dry-run]
       (a .csv works too)

What it does:
- Finds the header row and matches your columns to the template by name. Common variations work:
  "Employer" -> Company, "Position" or "Role" -> Job Title, "Applied On" -> Date Applied, "URL" -> Link, and so on.
- Standardizes statuses (for example "Phone screen" -> Interviewing, "Declined" -> Rejected) and keeps the
  original wording in Notes when it had to guess.
- Turns date text into real dates, keeps links (including =HYPERLINK formulas), and copies any columns the
  template doesn't have to the right of Notes, with their original headers, so nothing is lost.
- Never changes your original file.

Then run: python3 tools/import_tracker.py tracker/my_job_tracker.xlsx
Use --dry-run first to see how your columns will be matched. Needs openpyxl for .xlsx: pip3 install openpyxl
"""
import csv, datetime, re, sys, pathlib, shutil

ROOT = pathlib.Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "templates" / "job_tracker_template.xlsx"
HYPERLINK = re.compile(r'=HYPERLINK\(\s*"([^"]*)"\s*[,;]\s*"(.*)"\s*\)\s*$', re.I)

# template column -> header names people commonly use (lowercase, punctuation ignored)
FIELDS = {
    "Date Added":   ["date added", "added", "date found", "found", "date saved", "created", "date created", "date"],
    "Company":      ["company", "employer", "organization", "organisation", "company name", "business", "firm"],
    "Job Title":    ["job title", "title", "position", "role", "job", "job name", "position title", "role title"],
    "Location":     ["location", "city", "job location", "where", "office"],
    "Work Type":    ["work type", "remote", "remote hybrid", "remote hybrid onsite", "work arrangement", "arrangement",
                     "workplace", "workplace type", "work model", "on site remote"],
    "Salary Range": ["salary range", "salary", "pay", "pay range", "compensation", "comp", "salary pay", "base salary"],
    "Salary Min":   ["salary min", "min salary", "salary minimum", "minimum salary", "base min", "salary low", "low"],
    "Source":       ["source", "found on", "job board", "site", "platform", "where found", "board", "via"],
    "Link":         ["link", "url", "job link", "job url", "posting", "posting url", "posting link", "listing", "application link"],
    "Date Applied": ["date applied", "applied", "applied on", "application date", "applied date", "date of application", "submitted"],
    "Status":       ["status", "stage", "application status", "current status", "outcome", "result"],
    "Resume Match": ["resume match", "resume match %", "match", "match %", "match score"],
    "Fit Score":    ["fit score", "fit score %"],
    "Notes":        ["notes", "note", "comments", "comment", "next step", "next steps", "notes next step"],
}
TEMPLATE_COLS = ["Date Added", "Company", "Job Title", "Location", "Work Type", "Salary Range", "Salary Min", "Source",
                 "Link", "Date Applied", "Status", "Resume Match", "Fit Score", "Notes"]

STATUS_RULES = [  # first match wins; checked against the lowercase original status
    (r"offer|hired|accepted", "Offer"),
    (r"withdr|pulled|declined offer", "Withdrawn"),
    (r"second|2nd|final|onsite interview|panel|round 2|round two", "Second Round"),
    (r"interview|screen|phone|recruiter call|assessment|take.?home|hiring manager", "Interviewing"),
    (r"reject|declin|no longer|not selected|not moving|denied|closed|filled|ghost|no response|dead", "Rejected"),
    (r"did not apply|didn.t apply|not applied|skip|pass|not interested|not a fit", "Did Not Apply"),
    (r"applied|submitted|sent|in review|under review|pending", "Applied"),
    (r"saved|wish|bookmark|interested|to apply|todo|to do|planned|researching", "Saved"),
    (r"^new$|found|lead", "New"),
]
WORK_TYPE_RULES = [(r"remote", "Remote"), (r"hybrid", "Hybrid - Home Metro"), (r"on.?site|in.?office|office", "Onsite - Home Metro")]

def norm_header(h):
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(h or "").lower()).split())

def read_grid(path, sheet):
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [[(v, None) for v in row] for row in csv.reader(f)]
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [[(c.value, c.hyperlink.target if c.hyperlink else None) for c in row] for row in ws.iter_rows()]

def match_columns(headers):
    """Returns {template column: index} and the list of unmatched (index, header)."""
    normed = [norm_header(h) for h in headers]
    found = {}
    for field, names in FIELDS.items():          # exact matches first
        for i, h in enumerate(normed):
            if h and h in names and i not in found.values() and field not in found:
                found[field] = i
    for field, names in FIELDS.items():          # then "contains" matches, e.g. "Company Name (legal)"
        if field in found: continue
        for i, h in enumerate(normed):
            if h and i not in found.values() and any(re.search(r"\b" + re.escape(n) + r"\b", h) for n in names if len(n) > 3):
                found[field] = i; break
    extra = [(i, headers[i]) for i, h in enumerate(normed) if h and i not in found.values()]
    return found, extra

def to_date(v):
    if isinstance(v, datetime.datetime): return v
    if isinstance(v, datetime.date): return datetime.datetime(v.year, v.month, v.day)
    s = str(v or "").strip()
    if not s: return None
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%b %d, %Y", "%B %d, %Y", "%d %b %Y", "%m-%d-%Y", "%m/%d"):
        try:
            d = datetime.datetime.strptime(s, fmt)
            return d.replace(year=datetime.date.today().year) if fmt == "%m/%d" else d
        except ValueError:
            pass
    return None

def to_money(v):
    if isinstance(v, (int, float)): return int(v) if v >= 1000 else int(v * 1000)
    m = re.search(r"\$?\s*(\d[\d,]*(?:\.\d+)?)\s*([kK])?", str(v or ""))
    if not m: return None
    n = float(m.group(1).replace(",", ""))
    if m.group(2) or n < 1000: n *= 1000
    return int(n)

def map_status(raw, applied):
    s = str(raw or "").strip()
    if not s: return ("Applied" if applied else "New"), None
    for pat, val in STATUS_RULES:
        if re.search(pat, s.lower()): return val, (None if s.lower() == val.lower() else s)
    return ("Applied" if applied else "New"), s

def find_header_row(grid):
    for i, row in enumerate(grid[:15]):
        found, _ = match_columns([v for v, _ in row])
        if "Company" in found and "Job Title" in found: return i
    return None

def main():
    args = sys.argv[1:]
    if not args: sys.exit(__doc__)
    src = pathlib.Path(args[0]).expanduser()
    sheet = args[args.index("--sheet") + 1] if "--sheet" in args else None
    out = pathlib.Path(args[args.index("--out") + 1]).expanduser() if "--out" in args else ROOT / "tracker" / "my_job_tracker.xlsx"
    dry = "--dry-run" in args
    grid = read_grid(src, sheet)
    hi = find_header_row(grid)
    if hi is None:
        sys.exit("Couldn't find a header row with a company column and a job title column in the first 15 rows.\n"
                 "Rename those headers (for example 'Company' and 'Job Title') or pass --sheet with the right tab.")
    headers = [v for v, _ in grid[hi]]
    found, extra = match_columns(headers)
    print(f"Header row: {hi + 1}. Column matches:")
    for col in TEMPLATE_COLS:
        print(f"  {col:13} <- {headers[found[col]]!r}" if col in found else f"  {col:13} <- (none)")
    if extra: print("  Copied as extra columns: " + ", ".join(repr(h) for _, h in extra))
    rows = [r for r in grid[hi + 1:] if any(v not in (None, "") for v, _ in r)]
    if dry:
        print(f"\n{len(rows)} rows would be copied. Run again without --dry-run to create {out}."); return

    import openpyxl
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(TEMPLATE, out)
    wb = openpyxl.load_workbook(out); ws = wb["Applications"]
    for j, (_, h) in enumerate(extra):
        c = ws.cell(1, len(TEMPLATE_COLS) + 1 + j, str(h)); c._style = ws.cell(1, len(TEMPLATE_COLS))._style
        ws.column_dimensions[openpyxl.utils.get_column_letter(len(TEMPLATE_COLS) + 1 + j)].width = 18
    today = datetime.datetime.combine(datetime.date.today(), datetime.time())
    copied, guessed = 0, 0
    for r in rows:
        get = lambda col: (r[found[col]] if col in found and found[col] < len(r) else (None, None))
        company, _ = get("Company"); title, tlink = get("Job Title")
        tm = HYPERLINK.match(str(title or ""))
        if tm: tlink, title = tlink or tm.group(1), tm.group(2)
        if not str(company or "").strip() or not str(title or "").strip(): continue
        rowi = copied + 2
        link_v, link_h = get("Link")
        lm = HYPERLINK.match(str(link_v or ""))
        url = link_h or (lm.group(1) if lm else None) or (str(link_v) if str(link_v or "").startswith("http") else None) or \
              (tlink if str(tlink or "").startswith("http") else None)
        applied = to_date(get("Date Applied")[0])
        status, original = map_status(get("Status")[0], applied)
        notes = [str(get("Notes")[0]).strip()] if get("Notes")[0] not in (None, "") else []
        if original: notes.append(f"original status: {original}"); guessed += 1
        wt_raw = str(get("Work Type")[0] or "")
        wt = next((v for p, v in WORK_TYPE_RULES if re.search(p, wt_raw.lower())), None) if wt_raw else None
        if wt_raw and not wt: notes.append(f"work type: {wt_raw}")
        sal = get("Salary Range")[0]
        smin = to_money(get("Salary Min")[0]) if "Salary Min" in found else (to_money(sal) if sal not in (None, "") else None)
        values = {
            "Date Added": to_date(get("Date Added")[0]) or applied or today,
            "Company": str(company).strip(), "Job Title": str(title).strip(),
            "Location": get("Location")[0], "Work Type": wt, "Salary Range": sal, "Salary Min": smin,
            "Source": get("Source")[0], "Link": "Link" if url else None, "Date Applied": applied,
            "Status": status, "Resume Match": get("Resume Match")[0], "Fit Score": get("Fit Score")[0],
            "Notes": "; ".join(notes) or None}
        for ci, col in enumerate(TEMPLATE_COLS, 1):
            v = values[col]
            ws.cell(rowi, ci).value = v if not isinstance(v, str) else v.strip()
        if url: ws.cell(rowi, TEMPLATE_COLS.index("Link") + 1).hyperlink = url
        for j, (i, _) in enumerate(extra):
            v = r[i][0] if i < len(r) else None
            ws.cell(rowi, len(TEMPLATE_COLS) + 1 + j).value = v
        copied += 1
    wb.save(out)
    print(f"\nCopied {copied} jobs to {out}" + (f" ({guessed} statuses standardized; originals kept in Notes)" if guessed else ""))
    print(f"Your original file was not changed.\nNext: python3 tools/import_tracker.py {out}")

if __name__ == "__main__":
    main()
